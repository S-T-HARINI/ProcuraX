import csv
import io
from typing import List
from backend.app.models import DocumentPage, DocumentParseResponse


class DocumentService:
    @staticmethod
    def parse_document(file_bytes: bytes, filename: str) -> DocumentParseResponse:
        if not file_bytes or len(file_bytes.strip()) == 0:
            raise ValueError("Uploaded document is empty or contains only whitespace.")

        ext = filename.lower().split(".")[-1] if "." in filename else ""
        
        if ext == "pdf":
            return DocumentService._parse_pdf(file_bytes, filename)
        elif ext == "csv":
            return DocumentService._parse_csv(file_bytes, filename)
        elif ext in ["xlsx", "xls"]:
            return DocumentService._parse_excel(file_bytes, filename)
        else:
            raise ValueError(f"Unsupported file format '.{ext}'. Supported formats: PDF, CSV, XLSX.")

    @staticmethod
    def _parse_pdf(file_bytes: bytes, filename: str) -> DocumentParseResponse:
        pages: List[DocumentPage] = []
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))

            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    raise ValueError("PDF document is encrypted and password-protected.")

            if len(reader.pages) == 0:
                raise ValueError("PDF document contains no pages.")

            for i, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                clean_text = extracted.strip()
                pages.append(
                    DocumentPage(
                        page_number=i + 1,
                        text=clean_text if clean_text else f"[Page {i + 1}: No extractable text found]"
                    )
                )
        except ValueError:
            raise
        except Exception as err:
            raise ValueError(f"Failed to parse PDF document '{filename}': Corrupted or invalid format ({str(err)}).")

        raw_text = "\n\n".join([f"--- Page {p.page_number} ---\n{p.text}" for p in pages])
        return DocumentParseResponse(
            filename=filename,
            file_type="pdf",
            total_pages=len(pages),
            pages=pages,
            raw_text=raw_text,
        )

    @staticmethod
    def _parse_csv(file_bytes: bytes, filename: str) -> DocumentParseResponse:
        content = None
        for encoding in ["utf-8", "latin-1", "cp1252"]:
            try:
                content = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if content is None:
            raise ValueError(f"Failed to decode CSV document '{filename}' with standard text encodings.")

        lines = [line for line in content.splitlines() if line.strip()]
        if not lines:
            raise ValueError(f"CSV document '{filename}' contains no valid text lines.")

        # Validate CSV structure using csv module
        try:
            sample = "\n".join(lines[:10])
            reader = list(csv.reader(io.StringIO(sample)))
            if not reader or not any(row for row in reader):
                raise ValueError(f"CSV document '{filename}' contains no valid rows.")
        except Exception as err:
            raise ValueError(f"Malformed CSV content in '{filename}': {str(err)}")

        pages: List[DocumentPage] = []
        page_size = 50  # group 50 rows per logical page
        for i in range(0, len(lines), page_size):
            chunk_lines = lines[i:i + page_size]
            page_num = (i // page_size) + 1
            pages.append(DocumentPage(page_number=page_num, text="\n".join(chunk_lines)))

        return DocumentParseResponse(
            filename=filename,
            file_type="csv",
            total_pages=len(pages),
            pages=pages,
            raw_text=content,
        )

    @staticmethod
    def _parse_excel(file_bytes: bytes, filename: str) -> DocumentParseResponse:
        pages: List[DocumentPage] = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
            if not wb.sheetnames:
                raise ValueError(f"Excel workbook '{filename}' has no worksheets.")

            for i, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                rows_text = []
                for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    row_vals = [str(val).strip() if val is not None else "" for val in row]
                    if any(row_vals):
                        rows_text.append(f"Row {row_idx}: " + ", ".join(row_vals))
                
                sheet_content = f"Sheet: {sheet_name}\n" + ("\n".join(rows_text) if rows_text else "[Empty sheet]")
                pages.append(DocumentPage(page_number=i + 1, text=sheet_content))
        except ValueError:
            raise
        except Exception as err:
            raise ValueError(f"Failed to parse Excel document '{filename}': Corrupted or invalid format ({str(err)}).")

        raw_text = "\n\n".join([f"--- Page {p.page_number} ---\n{p.text}" for p in pages])
        return DocumentParseResponse(
            filename=filename,
            file_type="xlsx",
            total_pages=len(pages),
            pages=pages,
            raw_text=raw_text,
        )
